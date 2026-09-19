"""Resume import and extractive tailoring: never generate unsupported credentials."""
from io import BytesIO
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader

MAX_BYTES = 5 * 1024 * 1024

def extract(name, data):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Upload a nonempty file smaller than 5 MB.')
    try:
        if name.lower().endswith('.pdf'):
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 15:
                raise ValueError('Use an unencrypted PDF with at most 15 pages.')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        elif name.lower().endswith('.docx'):
            with ZipFile(BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('The expanded document is too large.')
            doc = Document(BytesIO(data))
            text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        elif name.lower().endswith('.txt'):
            text = data.decode('utf-8-sig')
        else:
            raise ValueError('Choose a PDF, DOCX, or UTF-8 TXT file.')
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('This document could not be read. Try another file or paste its text.') from exc
    if not 80 <= len(text.strip()) <= 40000:
        raise ValueError('The resume must contain 80–40,000 readable characters. For scanned PDFs, paste the text instead.')
    return text.strip()

def tokens(text):
    return set(re.findall(r'[a-z][a-z0-9+#.]{2,}', text.lower())) - {'the','and','with','for','you','our','are','will','from','your','that','this','have','work','team'}

def tailor(resume, description):
    """Reorder bullet runs within each original section; preserve every original line."""
    words = tokens(description)
    lines = resume.splitlines()
    output, run = [], []
    def flush():
        output.extend(sorted(run, key=lambda s: len(tokens(s) & words), reverse=True))
        run.clear()
    for line in lines:
        if re.match(r'^\s*[-•*▪]\s+', line):
            run.append(line)
        else:
            flush(); output.append(line)
    flush()
    shared = sorted(tokens(resume) & words)
    return '\n'.join(output), shared

def docx_bytes(text):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    style = doc.styles['Normal']
    style.font.name = 'Calibri'; style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(5)
    for i, line in enumerate(text.splitlines()):
        if i == 0:
            doc.add_heading(line, 0)
        elif line.strip() and (line.strip().isupper() or line.strip().rstrip(':').lower() in {'education', 'experience', 'skills', 'projects', 'summary', 'certifications'}):
            doc.add_heading(line, 2)
        else:
            doc.add_paragraph(line)
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()
